import openpyxl
from pathlib import Path
from docx import Document  # pip install python-docx


def read_messages():
    main_filepath = "data/IMapBook - CREW and discussions dataset.xlsx"
    records_file = Path(main_filepath)
    excel_obj = openpyxl.load_workbook(records_file)
    answer_sheet = excel_obj["CREW data"]

    messages = [i.value for i in answer_sheet['F'] if i.value != "Message"]
    return messages

def prepare_data():
    main_filepath = "data/IMapBook - CREW and discussions dataset.xlsx"
    records_file = Path(main_filepath)
    excel_obj = openpyxl.load_workbook(records_file)
    answer_sheet = excel_obj["CREW data"]

    messages = [i.value for i in answer_sheet['F'] if i.value != "Message"]
    classes = [i.value for i in answer_sheet['G'] if i.value != "CodePreliminary"]

    data = [['Id', 'Class', 'Text']]
    for i in range(len(messages)):
        data_item = [str(i), classes[i], "$" + messages[i] + "$"]
        data.append(data_item);

    return data


def read_classes():
    document = Document('../other/IMapBook - CREW codebook.docx')
    table = document.tables[0]
    data = []
    for i, row in enumerate(table.rows):
        text = (cell.text for cell in row.cells)
        if i == 0:
            continue
        row_data = tuple(text)
        if row_data[0] != '':
            data.append(row_data[0])
    return data
